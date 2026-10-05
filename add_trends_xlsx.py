import contextlib,io
with contextlib.redirect_stdout(io.StringIO()): exec(open('build3.py').read())
import json
R=json.load(open('output/dash3.json'))
gen=pd.DataFrame(R['genre'],columns=R['gnames'],index=R['months']).round(1); gen.insert(0,'Total music h',gen.sum(1).round(1))
gen['Average song year']=R['avgyear']; gen['Variety']=R['eff']; gen['Artists played']=R['arts']; gen['Top 5 share %']=R['top5']; gen['New songs %']=R['newtr']
gen=gen.reset_index().rename(columns={'index':'Month'})
eras=pd.DataFrame([dict(Era=i+1,Name=n,From=e['a'],To=e['b'],Hours=e['h'],**dict(zip(R['gnames'],[f"{v}%" for v in e['g']])),**{'Average song year':e['yr'],'Variety':e['eff'],'Top artists':', '.join(x[0] for x in e['top'][:5]),'Top songs':', '.join(x[0] for x in e['trk'])}) for i,(e,n) in enumerate(zip(R['eras'],['Pop and festival EDM','Eminem moves in','Hip-hop summer','The rock turn','Quieter, and instrumental','Deep classic rock','Metallica and Tuna','Wide open']))])
tg=pd.DataFrame([(a,{'P':'Pop','E':'Electronic & dance','H':'Hip-hop & R&B','I':'Israeli','A':'Alternative & modern rock','C':'Classic rock','M':'Hard rock & metal','N':'Instrumental & classical','O':'Everything else'}[t[0]],{'5':'1950s','6':'1960s','7':'1970s','8':'1980s','9':'1990s','0':'2000s','1':'2010s','2':'2020s','X':''}[t[1]],round(rank[a]/3600,1)) for a,t in zip(top400,T)],columns=['Artist','Genre','Decade','Hours'])
from openpyxl.styles import Font
with pd.ExcelWriter('output/Spotify_History_Organized.xlsx',engine='openpyxl',mode='a',if_sheet_exists='replace') as w:
    for name,d in [('Eras',eras),('Genre by Month',gen),('Artist Genres',tg)]:
        d.to_excel(w,sheet_name=name,index=False); ws=w.sheets[name]; ws.freeze_panes='B2'; ws.auto_filter.ref=ws.dimensions
        for col in ws.columns:
            L=max(len(str(c.value)) if c.value is not None else 0 for c in col[:200]); ws.column_dimensions[col[0].column_letter].width=min(max(9,L+2),40)
        for c in ws[1]: c.font=Font(bold=True)
from openpyxl import load_workbook; print(load_workbook('output/Spotify_History_Organized.xlsx').sheetnames)
