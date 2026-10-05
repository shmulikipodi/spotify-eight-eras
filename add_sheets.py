exec(open('build2.py').read().split("D={}")[0])
from openpyxl import load_workbook
out='output/Spotify_History_Organized.xlsx'
days_=pd.DataFrame({'Date':pd.date_range('2023-05-14','2026-09-27').date})
dd=df.groupby('date').agg(music=('sec',lambda s:0),total=('sec','sum'))
days_['Weekday']=pd.to_datetime(days_.Date).dt.day_name()
days_['Music min']=days_.Date.map(m.groupby('date').sec.sum()/60).fillna(0).round(0)
days_['Podcast min']=days_.Date.map(p.groupby('date').sec.sum()/60).fillna(0).round(0)
days_['Top artist']=days_.Date.map(m.groupby(['date','artist']).sec.sum().reset_index().sort_values('sec').groupby('date').tail(1).set_index('date').artist).fillna('')
s=df.groupby('sess').agg(Start=('ts','min'),End=('ts','max'),Songs=('sec','size'),sec=('sec','sum'),Top=('artist',lambda a:a.value_counts().index[0] if a.notna().any() else ''),First=('track','first')).reset_index(drop=True)
s['Minutes']=(s.sec/60).round(1); s['Start']=s.Start.dt.tz_localize(None); s['End']=s.End.dt.tz_localize(None)
s=s[['Start','End','Minutes','Songs','Top','First']].rename(columns={'Top':'Main artist','First':'First song'})
am=m.pivot_table(index='artist',columns='mon',values='sec',aggfunc='sum').fillna(0).div(3600).round(2)
am=am.loc[am.sum(axis=1).sort_values(ascending=False).index[:200]].reset_index().rename(columns={'artist':'Artist (hours per month)'})
g=m.groupby(['track','artist']).agg(Starts=('sec','size'),Finished=('reason_end',lambda x:round(100*(x=='trackdone').mean())),Skipped_10s=('sec',lambda x:round(100*(x<10).mean())),Back=('reason_start',lambda x:int((x=='backbtn').sum()))).reset_index()
g=g[g.Starts>=10].sort_values('Starts',ascending=False).rename(columns={'track':'Track','artist':'Artist','Finished':'Finished %','Skipped_10s':'Skipped <10s %','Back':'Times you went back'})
pe=p.groupby(['episode_show_name','episode_name']).agg(Minutes=('sec',lambda x:round(x.sum()/60)),First=('ts','min')).reset_index().sort_values('Minutes',ascending=False)
pe['First']=pe.First.dt.date; pe.columns=['Show','Episode','Minutes','First listened']
with pd.ExcelWriter(out,engine='openpyxl',mode='a',if_sheet_exists='replace') as w:
    for name,d in [('Every Day',days_),('Sessions',s),('Artists by Month',am),('Song Habits',g),('Podcast Episodes',pe)]:
        d.to_excel(w,sheet_name=name,index=False); ws=w.sheets[name]; ws.freeze_panes='B2' if name=='Artists by Month' else 'A2'; ws.auto_filter.ref=ws.dimensions
        for col in ws.columns:
            L=max(len(str(c.value)) if c.value is not None else 0 for c in col[:200]); ws.column_dimensions[col[0].column_letter].width=min(max(9,L+2),45)
        from openpyxl.styles import Font
        for c in ws[1]: c.font=Font(bold=True)
wb=load_workbook(out); print(wb.sheetnames)
